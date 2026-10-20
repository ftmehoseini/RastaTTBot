"""Shared workbook writer: Persian headers, RTL sheets, bold frozen header, filters, sized columns."""

import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY_FORMAT = "#,##0"
INT_FORMAT = "0"


@dataclass(frozen=True)
class Column:
    header: str
    width: int = 14
    number_format: str | None = None


def add_table_sheet(
    workbook: Workbook,
    title: str,
    columns: Sequence[Column],
    rows: Sequence[Sequence[Any]],
    totals: Sequence[Any] | None = None,
    first: bool = False,
) -> Worksheet:
    sheet = workbook.active if first else workbook.create_sheet()
    assert sheet is not None
    sheet.title = title[:31]
    sheet.sheet_view.rightToLeft = True
    sheet.append([c.header for c in columns])
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    for row in rows:
        sheet.append(list(row))
    if totals is not None:
        sheet.append(list(totals))
        for cell in sheet[sheet.max_row]:
            cell.font = TOTAL_FONT
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
        for cell in sheet[letter][1:]:
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if column.number_format:
                cell.number_format = column.number_format
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    return sheet


def new_workbook() -> Workbook:
    return Workbook()


def to_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def minutes_cell(minutes: int | None) -> str:
    """Durations as 'H:MM' text (Excel time formats break for negatives and >24h)."""
    if minutes is None:
        return ""
    sign = "-" if minutes < 0 else ""
    value = abs(minutes)
    return f"{sign}{value // 60}:{value % 60:02d}"
