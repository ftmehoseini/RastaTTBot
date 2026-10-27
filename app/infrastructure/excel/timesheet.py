"""The organization's monthly timesheet workbook, with its own formulas.

One sheet per employee and month, laid out like the organization's Excel file: a table of the
month's days (columns A–P) whose columns J–O are the workbook's formulas, and the summary block
under it. Only the inputs are written by the bot (weekday, date, check-in, check-out, paid leave,
notes); Excel calculates everything else when the file is opened, with exactly the formulas below.
`app/domain/attendance/timesheet.py` is the same arithmetic in Python for the bot's own reports.
"""

import io
import re
from collections.abc import Mapping, Sequence
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from app.application.dto.attendance import DayRow, EmployeeReport
from app.core.utils.jalali import iter_days, month_name, to_jalali, weekday_name
from app.domain.attendance.models import DayStatus
from app.domain.attendance.timesheet import SheetRules

# Header texts are referenced by the formulas, so they must stay exactly like this (note the double
# space in «مرخصی  با حقوق», as in the organization's file).
COLUMNS = [
    ("روز", 12),
    ("ماه", 10),
    ("تاریخ", 7),
    ("ساعت ورود", 11),
    ("شروع شناوری", 12),
    ("پایان شناوری", 12),
    ("ساعت مجاز خروج", 14),
    ("ساعت خروج", 11),
    ("مرخصی  با حقوق", 14),
    ("ساعت مفید", 11),
    ("کارکرد عادی روزانه", 17),
    ("تاخیر مجاز", 11),
    ("تاخیر غیر مجاز", 13),
    ("تعجیل", 9),
    ("اضافه کار روزانه", 16),
    ("توضیحات", 40),
]

# The workbook's calculated columns, verbatim; {t} is the table name.
FORMULAS = {
    "ساعت مفید": "{t}[[#This Row],[ساعت خروج]]-{t}[[#This Row],[ساعت ورود]]",
    "کارکرد عادی روزانه": (
        "IF(AND({t}[[#This Row],[ساعت ورود]]>0, {t}[[#This Row],[ساعت ورود]]>={t}[[#This Row],[شروع شناوری]]),"
        " IF({t}[[#This Row],[ساعت خروج]]<={t}[[#This Row],[ساعت مجاز خروج]],"
        " MAX({t}[[#This Row],[ساعت خروج]]-{t}[[#This Row],[ساعت ورود]], 0),"
        " {t}[[#This Row],[ساعت مجاز خروج]]-{t}[[#This Row],[ساعت ورود]]),"
        " IF({t}[[#This Row],[ساعت خروج]]<={t}[[#This Row],[ساعت مجاز خروج]],"
        " MAX({t}[[#This Row],[ساعت خروج]]-{t}[[#This Row],[شروع شناوری]], 0), "
        " {t}[[#This Row],[ساعت مجاز خروج]]-{t}[[#This Row],[شروع شناوری]]))"
    ),
    "تاخیر مجاز": (
        "IF(AND({t}[[#This Row],[ساعت ورود]]>={t}[[#This Row],[شروع شناوری]],"
        "+{t}[[#This Row],[ساعت ورود]]<={t}[[#This Row],[پایان شناوری]]),"
        "{t}[[#This Row],[ساعت ورود]]-{t}[[#This Row],[شروع شناوری]],"
        "IF({t}[[#This Row],[ساعت ورود]]>{t}[[#This Row],[پایان شناوری]],"
        " {t}[[#This Row],[پایان شناوری]]-{t}[[#This Row],[شروع شناوری]], 0))"
    ),
    "تاخیر غیر مجاز": (
        "IF({t}[[#This Row],[ساعت ورود]]>{t}[[#This Row],[پایان شناوری]],"
        " {t}[[#This Row],[ساعت ورود]]-{t}[[#This Row],[پایان شناوری]], 0)"
    ),
    "تعجیل": (
        "IF({t}[[#This Row],[ساعت خروج]]>0, MAX({t}[[#This Row],[ساعت مجاز خروج]]-{t}[[#This Row],[ساعت خروج]], 0), 0)"
    ),
    "اضافه کار روزانه": "MAX({t}[[#This Row],[ساعت مفید]]-{t}[[#This Row],[کارکرد عادی روزانه]], 0)",
}

TIME_FORMAT = "h:mm"
TOTAL_FORMAT = "[h]:mm:ss"
DAY_FORMAT = "[$-3000401]0"  # day number in Persian digits, as in the organization's file
TABLE_STYLE = "TableStyleMedium6"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "کارمند"
    title, n = base, 2
    while title in used:
        title, n = f"{base[:25]} {n}", n + 1
    used.add(title)
    return title


def _clock(minutes: int) -> time | timedelta:
    """Excel time value: a time of day, or more than a day for a check-out after midnight."""
    if 0 <= minutes < 1440:
        return time(minutes // 60, minutes % 60)
    return timedelta(minutes=minutes)


def _minutes_of(day: date, value: datetime) -> int:
    midnight = datetime.combine(day, time(0, 0), tzinfo=value.tzinfo)
    return int((value - midnight).total_seconds() // 60)


def _notes(row: DayRow | None, day: date, holidays: Mapping[date, str]) -> str:
    notes = []
    if day in holidays:
        notes.append(f"تعطیل رسمی: {holidays[day]}")
    if row is not None:
        if not row.is_workday and row.calculation.check_in is not None:
            notes.append("روز تعطیل: کل حضور اضافه‌کار")
        if row.leave_note:
            notes.append(row.leave_note)
        status = row.calculation.status
        if status == DayStatus.MISSING_CHECK_OUT:
            notes.append("خروج ثبت نشده")
        elif status == DayStatus.MISSING_CHECK_IN:
            notes.append("ورود ثبت نشده")
        if row.is_forgotten:
            notes.append("ثبت فراموشی")
        if row.is_manual_today:
            notes.append("ثبت دستی")
        if row.is_edited:
            notes.append("ویرایش‌شده")
    return " | ".join(notes)


def _write_summary(sheet: Worksheet, table: str, first: int, rules: SheetRules) -> None:
    """The summary block under the table, with the workbook's formulas."""
    r1, r2, r3, r4 = first, first + 1, first + 2, first + 3
    final = f"=B{r1}+MIN(K{r1}, P{r1})+B{r3}"
    if rules.surplus_factor:
        final += f"+P{r2}*{rules.surplus_factor:g}"
    if rules.penalty_factor:
        final += f"-P{r3}*{rules.penalty_factor:g}"
    cells = {
        f"A{r1}": "مجموع کارکرد عادی:",
        f"B{r1}": f"=SUM({table}[کارکرد عادی روزانه])",
        f"I{r1}": "مجموع اضافه کار:",
        f"K{r1}": f"= SUM({table}[اضافه کار روزانه])",
        f"O{r1}": "مقدار تاخیر و تعجیل قابل جبران:",
        f"P{r1}": f"=K{r2}+MIN(K{r3}, L{r3})+MIN(K{r4}, L{r4})",
        f"A{r2}": "مرخصی مجاز:",
        f"I{r2}": "مجموع تاخیر مجاز:",
        f"K{r2}": f"=SUM({table}[تاخیر مجاز])",
        f"O{r2}": "مقدار اضافه کار مازاد بر نیاز:",
        f"P{r2}": f"=MAX(K{r1}-P{r1},0)",
        f"A{r3}": "مرخصی مصرفی:",
        f"B{r3}": f"=SUM({table}[مرخصی  با حقوق])",
        f"I{r3}": "مجموع تاخیر غیرمجاز:",
        f"K{r3}": f"=SUM({table}[تاخیر غیر مجاز])",
        f"L{r3}": _clock(rules.late_cap),
        f"M{r3}": f"=MAX(K{r3}-L{r3},0)",
        f"O{r3}": "میزان تاخیر و تعجیل جبران نشده:",
        f"P{r3}": f"=MAX(P{r1}-K{r1},0)+MAX(K{r3}-L{r3},0)+MAX(K{r4}-L{r4},0)",
        f"A{r4}": "باقیمانده مرخصی:",
        f"B{r4}": f"=B{r2}-B{r3}",
        f"I{r4}": "مجموع تعجیل:",
        f"K{r4}": f"=SUM({table}[تعجیل])",
        f"L{r4}": _clock(rules.early_cap),
        f"M{r4}": f"=MAX(K{r4}-L{r4},0)",
        f"O{r4}": "کارکرد نهایی:",
        f"P{r4}": final,
    }
    for ref, value in cells.items():
        cell = sheet[ref]
        cell.value = value
        if isinstance(value, str) and not value.startswith("="):
            cell.font = Font(bold=True)
        elif ref[0] == "L":
            cell.number_format = TIME_FORMAT
        else:
            cell.number_format = TOTAL_FORMAT
    sheet[f"B{r2}"].number_format = TOTAL_FORMAT  # مرخصی مجاز: filled in by the admin


def add_timesheet_sheet(
    workbook: Workbook,
    report: EmployeeReport,
    rules: SheetRules,
    holidays: Mapping[date, str],
    table_name: str,
    title: str,
) -> Worksheet:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "D2"
    for col, (header, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, col, header)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = width
    rows = {r.work_date: r for r in report.rows}
    days = iter_days(report.start, report.end)
    month = month_name(to_jalali(report.start).month)
    for index, day in enumerate(days, start=2):
        row = rows.get(day)
        c = row.calculation if row else None
        values: dict[str, object] = {
            "روز": weekday_name(day),
            "ماه": month,
            "تاریخ": to_jalali(day).day,
            "شروع شناوری": _clock(rules.floating_start),
            "پایان شناوری": _clock(rules.floating_end),
            "ساعت مجاز خروج": _clock(rules.allowed_exit),
            "توضیحات": _notes(row, day, holidays) or None,
        }
        # A check-out without a check-in is left out: the formulas would read it as presence from 0:00.
        if c is not None and c.check_in is not None:
            check_in = _minutes_of(day, c.check_in)
            values["ساعت ورود"] = _clock(check_in)
            if c.check_out is not None:
                values["ساعت خروج"] = _clock(_minutes_of(day, c.check_out))
            if not row.is_workday:
                # Day off: E = F = G = the check-in makes the same formulas count all presence as
                # overtime (no normal work, lateness or early departure), as the bot does.
                for header in ("شروع شناوری", "پایان شناوری", "ساعت مجاز خروج"):
                    values[header] = _clock(check_in)
        if row is not None and row.paid_leave_minutes:
            values["مرخصی  با حقوق"] = _clock(row.paid_leave_minutes)
        for col, (header, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(index, col)
            if header in FORMULAS:
                cell.value = "=" + FORMULAS[header].format(t=table_name)
            else:
                cell.value = values.get(header)
            if header == "تاریخ":
                cell.number_format = DAY_FORMAT
            elif 4 <= col <= 15:
                cell.number_format = TIME_FORMAT
    last = len(days) + 1
    table = Table(displayName=table_name, ref=f"A1:P{last}")
    table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
    # Declaring the calculated columns lets Excel fill the formulas into rows added by hand.
    table.tableColumns = [
        TableColumn(
            id=i,
            name=header,
            calculatedColumnFormula=(
                TableFormula(attr_text=FORMULAS[header].format(t=table_name)) if header in FORMULAS else None
            ),
        )
        for i, (header, _) in enumerate(COLUMNS, start=1)
    ]
    sheet.add_table(table)
    _write_summary(sheet, table_name, last + 1, rules)
    return sheet


def build_timesheet_workbook(
    reports: Sequence[EmployeeReport], rules: SheetRules, holidays: Mapping[date, str]
) -> bytes:
    """One timesheet sheet per (full-time) employee; each report must cover one Jalali month."""
    workbook = Workbook()
    default = workbook.active
    used: set[str] = set()
    month = month_name(to_jalali(reports[0].start).month) if reports else "ماه"
    for i, report in enumerate((r for r in reports if not r.hourly), start=1):
        add_timesheet_sheet(workbook, report, rules, holidays, f"{month}_{i}", _sheet_title(report.full_name, used))
    if len(workbook.worksheets) > 1:
        workbook.remove(default)
    else:
        default.title = "بدون داده"
        default.sheet_view.rightToLeft = True
        default["A1"] = "کارمند تمام‌وقتی در این گزارش نیست."
    workbook.calculation.fullCalcOnLoad = True  # Excel and LibreOffice recalculate on open
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
