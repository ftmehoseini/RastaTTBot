"""Staff list (roster) workbook: the file the admin uploads, and the current list as a download."""

import io
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.core.utils.text import clean_text, normalize_employee_code, to_persian_digits
from app.infrastructure.excel.styles import Column, add_table_sheet, new_workbook, to_bytes

HEADERS = ("شماره موبایل یا پرسنلی", "نام و نام خانوادگی", "نوع همکاری")
MAX_ERRORS = 20


@dataclass(frozen=True)
class RosterRow:
    code: str
    full_name: str | None
    hourly: bool


@dataclass
class RosterParse:
    rows: list[RosterRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


def _column_roles(header: Sequence[object]) -> dict[str, int] | None:
    """Which column holds the number, the name and the type, from a header row; None if not a header."""
    roles: dict[str, int] = {}
    for index, value in enumerate(header):
        text = str(value or "").strip().lower()
        if not text:
            continue
        if "code" not in roles and any(k in text for k in ("شماره", "موبایل", "تلفن", "کد", "mobile", "phone", "code")):
            roles["code"] = index
        elif "type" not in roles and any(k in text for k in ("نوع", "همکاری", "type")):
            roles["type"] = index
        elif "name" not in roles and any(k in text for k in ("نام", "name")):
            roles["name"] = index
    return roles if "code" in roles else None


def parse_roster(content: bytes, max_name_length: int = 120) -> RosterParse:
    result = RosterParse()
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        result.errors.append("فایل اکسل معتبر نیست. لطفاً فایل ‎.xlsx‎ بفرستید.")
        return result
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    roles: dict[str, int] | None = None
    start = 0
    for i, row in enumerate(rows[:5]):
        roles = _column_roles(row)
        if roles:
            start = i + 1
            break
    if roles is None:  # no header: number, name, type in the first three columns
        roles = {"code": 0, "name": 1, "type": 2}
    seen: dict[str, int] = {}
    for number, row in enumerate(rows[start:], start=start + 1):

        def cell(role: str, row: tuple = row) -> object:
            index = roles.get(role)
            return row[index] if index is not None and index < len(row) else None

        raw = cell("code")
        if raw is None or str(raw).strip() == "":
            if any(v not in (None, "") for v in row):
                result.errors.append(f"ردیف {to_persian_digits(number)}: شماره خالی است.")
            continue
        code = normalize_employee_code(str(raw))
        if code is None:
            result.errors.append(f"ردیف {to_persian_digits(number)}: «{raw}» شماره‌ی معتبری نیست.")
            continue
        if code in seen:
            result.errors.append(
                f"ردیف {to_persian_digits(number)}: شماره {to_persian_digits(code)} تکراری است "
                f"(ردیف {to_persian_digits(seen[code])})."
            )
            continue
        seen[code] = number
        name = clean_text(str(cell("name") or ""), max_name_length) or None
        kind = str(cell("type") or "").strip().lower()
        result.rows.append(RosterRow(code, name, hourly="ساعتی" in kind or "hourly" in kind))
    if not result.rows and not result.errors:
        result.errors.append("هیچ شماره‌ای در فایل پیدا نشد.")
    result.errors = result.errors[:MAX_ERRORS]
    return result


def build_roster_workbook(rows: Sequence[tuple[str, str | None, bool, str | None]]) -> bytes:
    """(code, list name, hourly, registered user's name or None) per entry; also the upload template."""
    workbook = new_workbook()
    columns = [Column(HEADERS[0], 24, "@"), Column(HEADERS[1], 26), Column(HEADERS[2], 14), Column("وضعیت ثبت‌نام", 30)]
    data = [
        [code, name or "", "ساعتی" if hourly else "تمام‌وقت", f"ثبت‌نام کرده: {user}" if user else "هنوز ثبت‌نام نکرده"]
        for code, name, hourly, user in rows
    ]
    add_table_sheet(workbook, "فهرست کارکنان", columns, data, first=True)
    guide = workbook.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    for line in (
        "ستون اول: شماره موبایل (مثل 09121234567) یا شماره پرسنلی — همان چیزی که کارمند هنگام ثبت‌نام می‌نویسد.",
        "ستون دوم: نام و نام خانوادگی (اختیاری). اگر پر باشد، کارمند با همین نام ثبت می‌شود.",
        "ستون سوم: نوع همکاری (اختیاری): «تمام‌وقت» یا «ساعتی».",
        "ستون «وضعیت ثبت‌نام» فقط برای دیدن است و هنگام بارگذاری نادیده گرفته می‌شود.",
        "با بارگذاری فایل جدید، شماره‌هایی که در فایل نیستند و هنوز ثبت‌نام نکرده‌اند حذف می‌شوند.",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    return to_bytes(workbook)
