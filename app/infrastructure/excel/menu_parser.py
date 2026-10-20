"""Weekly food menu template and parser. All validation messages are Persian for the uploader."""

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.utils.jalali import JalaliDate, format_date_numeric, to_jalali, week_range, weekday_name
from app.core.utils.text import normalize_digits, to_persian_digits
from app.infrastructure.excel.styles import Column, add_table_sheet, new_workbook, to_bytes

REQUIRED_HEADERS = ("jalali_date", "food_name", "price_toman")
OPTIONAL_HEADERS = ("weekday_name", "is_active")

HEADER_ALIASES: dict[str, str] = {
    "jalali_date": "jalali_date",
    "تاریخ": "jalali_date",
    "تاریخ شمسی": "jalali_date",
    "weekday_name": "weekday_name",
    "روز": "weekday_name",
    "روز هفته": "weekday_name",
    "food_name": "food_name",
    "نام غذا": "food_name",
    "غذا": "food_name",
    "price_toman": "price_toman",
    "قیمت": "price_toman",
    "قیمت (تومان)": "price_toman",
    "قیمت تومان": "price_toman",
    "is_active": "is_active",
    "فعال": "is_active",
}

_TRUE = {"1", "true", "yes", "y", "بله", "فعال", "✓"}
_FALSE = {"0", "false", "no", "n", "خیر", "غیرفعال"}
MAX_ROWS = 300
MAX_ERRORS_SHOWN = 15


@dataclass(frozen=True)
class MenuRow:
    row_number: int
    menu_date: date
    food_name: str
    price_toman: int
    is_active: bool


@dataclass
class MenuParseResult:
    rows: list[MenuRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    week_start: date | None = None

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


def normalize_persian(text: str) -> str:
    """Unify Arabic/Persian letter variants and spacing so 'سه شنبه' == 'سه‌شنبه'."""
    return text.replace("ي", "ی").replace("ك", "ک").replace("‌", "").replace(" ", "").strip()


def normalize_food_name(text: str) -> str:
    return " ".join(text.replace("ي", "ی").replace("ك", "ک").split())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_digits(str(value)).strip()


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()  # a real Excel date cell is Gregorian
    if isinstance(value, date):
        return value
    return JalaliDate.parse(_cell_text(value)).to_gregorian()


def _parse_price(value: Any, max_price: int) -> int:
    text = _cell_text(value).replace(",", "").replace("٬", "").replace("،", "")
    if not text.isdigit():
        raise ValueError
    price = int(text)
    if price <= 0 or price > max_price:
        raise ValueError
    return price


def parse_menu(content: bytes, today: date, max_price: int) -> MenuParseResult:
    result = MenuParseResult()
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError):
        result.errors.append("فایل اکسل معتبر نیست. لطفاً فایل ‎.xlsx‎ را از روی قالب بسازید.")
        return result

    sheet = workbook.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        result.errors.append("فایل خالی است.")
        return result

    columns: dict[str, int] = {}
    for index, raw in enumerate(header_row):
        key = HEADER_ALIASES.get(_cell_text(raw).lower()) or HEADER_ALIASES.get(_cell_text(raw))
        if key and key not in columns:
            columns[key] = index
    missing = [h for h in REQUIRED_HEADERS if h not in columns]
    if missing:
        result.errors.append(f"ستون‌های لازم پیدا نشد: {', '.join(missing)}. از قالب دانلودی استفاده کنید.")
        return result

    def get(row: tuple, key: str) -> Any:
        index = columns.get(key)
        return row[index] if index is not None and index < len(row) else None

    seen: set[tuple[date, str]] = set()
    for row_number, row in enumerate(rows_iter, start=2):
        if row_number > MAX_ROWS + 1:
            result.errors.append(f"تعداد ردیف‌ها بیشتر از {to_persian_digits(MAX_ROWS)} است.")
            break
        if all(_cell_text(v) == "" for v in row):
            continue
        label = f"ردیف {to_persian_digits(row_number)}"
        try:
            menu_date = _parse_date(get(row, "jalali_date"))
        except ValueError:
            result.errors.append(f"{label}: تاریخ نامعتبر است (نمونه‌ی درست: ۱۴۰۵/۰۷/۲۶).")
            continue

        weekday_cell = _cell_text(get(row, "weekday_name"))
        if weekday_cell and normalize_persian(weekday_cell) != normalize_persian(weekday_name(menu_date)):
            result.errors.append(
                f"{label}: روز «{weekday_cell}» با تاریخ {format_date_numeric(menu_date)} "
                f"({weekday_name(menu_date)}) هم‌خوانی ندارد."
            )

        food_name = normalize_food_name(_cell_text(get(row, "food_name")))
        if not food_name:
            result.errors.append(f"{label}: نام غذا خالی است.")
            continue
        if len(food_name) > 120:
            result.errors.append(f"{label}: نام غذا بیش از حد طولانی است.")
            continue

        try:
            price = _parse_price(get(row, "price_toman"), max_price)
        except ValueError:
            result.errors.append(f"{label}: قیمت باید عدد صحیح مثبت (به تومان) باشد.")
            continue

        active_cell = _cell_text(get(row, "is_active")).lower()
        if active_cell in ("", *_TRUE):
            is_active = True
        elif active_cell in _FALSE:
            is_active = False
        else:
            result.errors.append(f"{label}: مقدار is_active باید ۱ یا ۰ باشد.")
            continue

        key = (menu_date, food_name)
        if key in seen:
            result.errors.append(f"{label}: «{food_name}» برای این تاریخ تکراری است.")
            continue
        seen.add(key)
        result.rows.append(MenuRow(row_number, menu_date, food_name, price, is_active))

    workbook.close()
    if not result.rows and not result.errors:
        result.errors.append("هیچ ردیف غذایی در فایل پیدا نشد.")
    if result.rows:
        weeks = {week_range(r.menu_date)[0] for r in result.rows}
        if len(weeks) > 1:
            result.errors.append("همه‌ی ردیف‌ها باید مربوط به یک هفته (شنبه تا جمعه) باشند.")
        else:
            week_start = weeks.pop()
            if week_start + timedelta(days=6) < today:
                result.errors.append("این منو برای هفته‌ای است که تمام شده است.")
            result.week_start = week_start
    if len(result.errors) > MAX_ERRORS_SHOWN:
        extra = len(result.errors) - MAX_ERRORS_SHOWN
        result.errors = [*result.errors[:MAX_ERRORS_SHOWN], f"… و {to_persian_digits(extra)} خطای دیگر."]
    return result


def build_menu_template(week_start: date) -> bytes:
    """A ready-to-fill template for the given week, one example row per day (Saturday–Thursday)."""
    workbook = new_workbook()
    columns = [
        Column("jalali_date", 14),
        Column("weekday_name", 12),
        Column("food_name", 28),
        Column("price_toman", 14, "#,##0"),
        Column("is_active", 10),
    ]
    rows = []
    for offset in range(6):
        day = week_start + timedelta(days=offset)
        rows.append([to_jalali(day).iso(), weekday_name(day), "", None, 1])
    sheet = add_table_sheet(workbook, "menu", columns, rows, first=True)
    validation = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"E2:E{MAX_ROWS + 1}")

    guide = workbook.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    for line in (
        "راهنمای پر کردن منوی هفتگی",
        "jalali_date: تاریخ شمسی به شکل ۱۴۰۵/۰۷/۲۶ (متنی). همه‌ی ردیف‌ها باید در یک هفته باشند.",
        "weekday_name: نام روز هفته (اختیاری، برای کنترل اشتباه تاریخ).",
        "food_name: نام غذا. برای چند غذا در یک روز، چند ردیف با همان تاریخ بنویسید.",
        "price_toman: قیمت به تومان، عدد صحیح.",
        "is_active: ۱ یعنی قابل سفارش، ۰ یعنی نمایش داده نشود.",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 90
    return to_bytes(workbook)
