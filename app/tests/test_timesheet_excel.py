"""The timesheet workbook: its formulas, recalculated by LibreOffice, give the bot's numbers."""

import asyncio
import datetime as dt
import io
import os
import shutil
import subprocess
from datetime import date, time
from pathlib import Path

import openpyxl
import pytest

from app.application.dto.attendance import CorrectionInput
from app.domain.attendance.models import CorrectionType
from app.infrastructure.excel.timesheet import FORMULAS, build_timesheet_workbook
from app.tests.conftest import tehran

DAYS = (
    (date(2026, 10, 17), (8, 38), (17, 6)),
    (date(2026, 10, 18), (11, 34), (19, 0)),
    (date(2026, 10, 19), (9, 25), (16, 27)),
    (date(2026, 10, 20), (7, 40), (16, 0)),
)


async def _month_workbook(services, make_user, clock) -> tuple[bytes, object]:
    user = await make_user("مریم کریمی")
    clock.set(tehran(2026, 10, 21, 20, 0))
    for day, check_in, check_out in DAYS:
        await services.attendance.submit_correction(
            user, CorrectionInput(CorrectionType.BOTH, day, time(*check_in), time(*check_out))
        )
    report = await services.attendance.my_month(user)
    content = build_timesheet_workbook([report], services.attendance.sheet_rules, set())
    return content, report


async def test_workbook_layout_and_formulas(services, make_user, clock) -> None:
    content, _ = await _month_workbook(services, make_user, clock)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["مریم کریمی"]
    table = next(iter(ws.tables.values()))
    assert table.ref == "A1:P31"  # مهر has 30 days
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["روز", "ماه", "تاریخ", "ساعت ورود"] and headers[8] == "مرخصی  با حقوق"
    # 17/10/2026 = 25 مهر: inputs written, calculated columns are the workbook's formulas
    row = 26
    assert (ws.cell(row, 1).value, ws.cell(row, 2).value, ws.cell(row, 3).value) == ("شنبه", "مهر", 25)
    assert ws.cell(row, 4).value == time(8, 38) and ws.cell(row, 8).value == time(17, 6)
    assert ws.cell(row, 11).value == "=" + FORMULAS["کارکرد عادی روزانه"].format(t=table.displayName)
    assert ws["P35"].value == "=B32+MIN(K32, P32)+B34+P33*0.8-P34*1.6"
    assert ws["L34"].value == time(7, 0) and ws["L35"].value == time(3, 0)
    assert "ثبت فراموشی" in ws.cell(row, 16).value


def _minutes(value) -> int:
    if isinstance(value, dt.time):
        return value.hour * 60 + value.minute + round(value.second / 60)
    if isinstance(value, dt.timedelta):
        return round(value.total_seconds() / 60)
    if isinstance(value, dt.datetime):  # LibreOffice may store a duration as a date-time
        return round((value - dt.datetime(1899, 12, 30)).total_seconds() / 60)
    return round((value or 0) * 1440)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
async def test_libreoffice_recalculation_matches_the_bot(services, make_user, clock, tmp_path: Path) -> None:
    content, report = await _month_workbook(services, make_user, clock)
    source = tmp_path / "in" / "timesheet.xlsx"
    source.parent.mkdir()
    source.write_bytes(content)
    out = tmp_path / "out"
    env = {**os.environ, "HOME": str(tmp_path)}
    await asyncio.to_thread(
        subprocess.run,
        ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(source)],
        check=True,
        capture_output=True,
        env=env,
        timeout=120,
    )
    ws = openpyxl.load_workbook(out / "timesheet.xlsx", data_only=True).active
    rows = {r.work_date: r for r in report.rows}
    for index, (day, _, _) in enumerate(DAYS):
        excel_row = 26 + index
        c = rows[day].calculation
        got = [_minutes(ws.cell(excel_row, col).value) for col in range(10, 16)]
        want = [
            c.gross_worked_minutes,
            c.normal_minutes,
            c.floating_late_minutes,
            c.late_minutes,
            c.early_departure_minutes,
            c.overtime_minutes,
        ]
        assert got == want, day
    m = report.sheet
    summary = {
        "B32": m.normal,
        "K32": m.overtime,
        "P32": m.compensable,
        "K33": m.floating_late,
        "P33": m.overtime_surplus,
        "K34": m.late,
        "P34": m.uncompensated,
        "K35": m.early,
        "P35": m.final_work,
    }
    assert {k: _minutes(ws[k].value) for k in summary} == summary
