from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.application.use_cases.food_broadcast import broadcast_menu
from app.core.config.policy import load_policy
from app.core.exceptions import ConflictError, PermissionDeniedError, ValidationError
from app.core.logging.setup import redact
from app.core.security.rbac import Permission, RoleCode, has_permission, is_admin, permissions_for
from app.domain.audit.models import AuditAction, AuditLog
from app.domain.notifications.models import NotificationType
from app.infrastructure.excel.menu_parser import parse_menu
from app.infrastructure.excel.reports import build_attendance_workbook
from app.infrastructure.integrations.messenger import split_text
from app.tests.conftest import tehran
from app.tests.test_food import VALID_ROWS, menu_file


def test_role_permission_matrix() -> None:
    assert has_permission([RoleCode.EMPLOYEE], Permission.ATTENDANCE_SELF)
    assert not has_permission([RoleCode.EMPLOYEE], Permission.ATTENDANCE_REPORT_ALL)
    assert has_permission([RoleCode.ADMIN_1_SECRETARY], Permission.LEAVE_APPROVE_STEP_1)
    assert not has_permission([RoleCode.ADMIN_1_SECRETARY], Permission.LEAVE_APPROVE_STEP_2)
    assert has_permission([RoleCode.ADMIN_2_MANAGER], Permission.LEAVE_APPROVE_STEP_2)
    assert not has_permission([RoleCode.ADMIN_2_MANAGER], Permission.FOOD_MENU_MANAGE)
    assert not has_permission([RoleCode.ADMIN_1_SECRETARY], Permission.USERS_MANAGE_ROLES)
    assert permissions_for([RoleCode.SUPER_ADMIN]) == frozenset(Permission)
    # Multiple roles combine.
    combined = permissions_for([RoleCode.EMPLOYEE, RoleCode.ADMIN_1_SECRETARY])
    assert Permission.FOOD_ORDER in combined and Permission.FOOD_BROADCAST in combined
    assert is_admin([RoleCode.EMPLOYEE, RoleCode.ADMIN_2_MANAGER]) and not is_admin([RoleCode.EMPLOYEE])


async def test_employee_cannot_see_team_reports(services, make_user) -> None:
    employee = await make_user()
    with pytest.raises(PermissionDeniedError):
        await services.attendance.team_reports(employee, date(2026, 10, 17), date(2026, 10, 19))


async def test_only_super_admin_manages_roles_and_it_is_audited(services, make_user) -> None:
    super_admin = await make_user("مدیر کل", RoleCode.SUPER_ADMIN)
    secretary = await make_user("منشی", RoleCode.ADMIN_1_SECRETARY)
    target = await make_user("کارمند")
    with pytest.raises(PermissionDeniedError):
        await services.users.set_role(secretary, target.id, RoleCode.ADMIN_2_MANAGER, True)

    updated = await services.users.set_role(super_admin, target.id, RoleCode.ADMIN_2_MANAGER, True)
    assert RoleCode.ADMIN_2_MANAGER in updated.roles
    updated = await services.users.set_role(super_admin, target.id, RoleCode.ADMIN_2_MANAGER, False)
    assert RoleCode.ADMIN_2_MANAGER not in updated.roles
    with pytest.raises(ValidationError):  # cannot lock yourself out
        await services.users.set_role(super_admin, super_admin.id, RoleCode.SUPER_ADMIN, False)

    actions = (
        (await services.session.execute(select(AuditLog.action).where(AuditLog.actor_user_id == super_admin.id)))
        .scalars()
        .all()
    )
    assert actions.count(AuditAction.ROLE_ASSIGNED) == 1 and actions.count(AuditAction.ROLE_REVOKED) == 1


async def test_registration_and_bootstrap_admin(services, policy) -> None:
    services._bootstrap_ids = frozenset({777})
    admin = await services.users.register(777, 777, "مدیر سیستم", "boss")
    assert RoleCode.SUPER_ADMIN in admin.roles and RoleCode.EMPLOYEE in admin.roles
    employee = await services.users.register(778, 778, "زهرا احمدی", None)
    assert employee.roles == {RoleCode.EMPLOYEE}
    with pytest.raises(ConflictError):
        await services.users.register(778, 778, "زهرا احمدی", None)
    for bad in ("ab", "/start", "Ali 123"):
        with pytest.raises(ValidationError):
            services.users.validate_full_name(bad)


async def test_require_activation_policy(services, policy) -> None:
    policy.users.require_admin_activation = True
    user = await services.users.register(900, 900, "کاربر جدید", None)
    assert not user.is_active


async def test_notification_dedup_and_bulk_resume(services, make_user, messenger, clock) -> None:
    secretary = await make_user("منشی", RoleCode.ADMIN_1_SECRETARY)
    a = await make_user("آرش")
    b = await make_user("بهار")
    clock.set(tehran(2026, 10, 18, 8, 0))
    parsed = parse_menu(menu_file(VALID_ROWS), date(2026, 10, 18), 100_000_000)
    menu = (await services.food.import_menu(secretary, parsed, "m.xlsx", None)).menu

    messenger.fail_chats.add(b.chat_id)
    first = await broadcast_menu(services, secretary.id, menu.id)
    assert (first.sent, first.failed) == (2, 1)  # secretary is also an employee
    messenger.fail_chats.clear()
    second = await broadcast_menu(services, secretary.id, menu.id)
    assert (second.sent, second.skipped) == (1, 2)  # only the failed one is retried
    assert len(messenger.texts_to(a.chat_id)) == 1
    assert len(messenger.texts_to(b.chat_id)) == 1
    outcome = await services.notifications.notify_user(a, "x", NotificationType.GENERIC, dedup_key="k1")
    assert outcome == "SENT"
    assert await services.notifications.notify_user(a, "x", NotificationType.GENERIC, dedup_key="k1") == "SKIPPED"


def test_split_text_respects_limit() -> None:
    text = "\n".join(f"line {i} " + "x" * 50 for i in range(300))
    chunks = split_text(text, 1000)
    assert all(len(c) <= 1000 for c in chunks)
    assert "\n".join(chunks) == text
    assert split_text("short") == ["short"]
    assert all(len(c) <= 100 for c in split_text("y" * 450, 100))


def test_policy_overrides_and_validation(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[2] / "config"
    (tmp_path / "default.yaml").write_text((root / "default.yaml").read_text(encoding="utf-8"), encoding="utf-8")
    (tmp_path / "production.yaml").write_text("food:\n  ordering_deadline_time: '09:30'\n", encoding="utf-8")
    policy = load_policy(tmp_path, "production")
    assert policy.food.ordering_deadline_time.strftime("%H:%M") == "09:30"
    assert policy.attendance.standard_workday_minutes == 528  # untouched default
    (tmp_path / "broken.yaml").write_text("attendance:\n  workdays: [sat, xyz]\n", encoding="utf-8")
    with pytest.raises(ValueError):
        load_policy(tmp_path, "broken")


def test_tokens_are_redacted() -> None:
    assert "<redacted-token>" in redact(
        "https://api.telegram.org/bot123456789:AAHdqTcvCH1vGWJxfSeofSAs0K5PALDsaw/getMe"
    )


async def test_attendance_excel_has_rtl_header_and_filter(services, make_user, clock) -> None:
    user = await make_user("علی")
    clock.set(tehran(2026, 10, 19, 12, 0))
    report = await services.attendance.my_week(user)
    import io

    workbook = load_workbook(io.BytesIO(build_attendance_workbook([report])))
    sheet = workbook.worksheets[0]
    assert sheet.sheet_view.rightToLeft
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].value == "نام کارمند" and sheet["A1"].font.bold
    assert sheet.max_row == 1 + len(report.rows)
